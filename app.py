from flask import Flask, render_template, request, send_file
from docx import Document
import os

app = Flask(__name__)

# Function to generate bail application
def generate_bail_application(data, template_path, output_path):
    doc = Document(template_path)

    for paragraph in doc.paragraphs:
        full_text = paragraph.text
        for key, value in data.items():
            if key in full_text:
                full_text = full_text.replace(key, value)  # Replace placeholders
        paragraph.text = full_text  # Update paragraph text

    doc.save(output_path)

@app.route("/", methods=["GET", "POST"])
def form():
    if request.method == "POST":
        # Get form data
        court_name = request.form["court_name"]
        case_number = request.form["case_number"]
        applicant_name = request.form["applicant_name"]
        fir_details = request.form["fir_details"]
        police_station = request.form["police_station"]
        facts_of_fir = request.form["facts_of_fir"]  

        # ✅ Fetch multiple grounds for bail correctly
        grounds_for_bail = request.form.getlist("grounds_for_bail[]")  
        
        # ✅ Number each ground properly
        grounds_text = "\n".join([f"{i+1}. {ground}" for i, ground in enumerate(grounds_for_bail) if ground.strip()])

        under_sections = request.form["under_sections"]
        date = request.form["date"]
        advocate_for = request.form["advocate_for"]

        # Data mapping for placeholders
        data = {
            "{COURT_NAME}": court_name,
            "{CASE_NUMBER}": case_number,
            "{APPLICANT_NAME}": applicant_name,
            "{FIR_DETAILS}": fir_details,
            "{UNDER_SECTIONS}": under_sections,
            "{FACTS_OF_FIR}": facts_of_fir,
            "{GROUNDS_FOR_BAIL}": grounds_text,  # ✅ Properly formatted numbered inputs
            "{POLICE_STATION}": police_station,
            "{DATE}": date,
            "{ADVOCATE_FOR}": advocate_for
        }

        # File paths
        template_file = "bail_template.docx"
        output_file = "generated_bail_application.docx"

        # Ensure template exists
        if not os.path.exists(template_file):
            return "Error: Template file not found!", 404

        # Generate document
        generate_bail_application(data, template_file, output_file)

        # Send the file for download
        return send_file(output_file, as_attachment=True, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    return render_template("form.html")

if __name__ == "__main__":
    app.run(debug=True)
